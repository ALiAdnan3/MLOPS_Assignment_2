"""Build the final submission report from reports/REPORT_FINAL.md + screenshots/.

  python reports/build_final_report.py   ->  reports/23L0839_VisionGuard.docx  (then word_finalize.py -> PDF)

Source markers:
  ![label=name; caption](path)   figure, numbered automatically; {fig:name} in the text becomes "Figure N"
  [[OPT file.png | caption]]      optional screenshot: included only if screenshots/file.png exists
  CNN_TABLE / TOY_TABLE / APPLY_TABLE   tables generated from the CSV results
"""
import csv
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_report import add_code, add_inline, add_table, page_break, title_page  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "reports" / "REPORT_FINAL.md"
OUT = ROOT / "reports" / "23L0839_VisionGuard.docx"
SHOTS = ROOT / "screenshots"


def md_table(rows, header):
    lines = ["| " + " | ".join(header) + " |", "|" + "---|" * len(header)]
    lines += ["| " + " | ".join(str(c) for c in r) + " |" for r in rows]
    return "\n".join(lines)


def tables():
    cnn = list(csv.DictReader(open(ROOT / "reports" / "cnn_layers.csv")))
    cnn_t = md_table([(r["layer"], r["type"], r["output_shape"], f"{int(r['params']):,}") for r in cnn],
                     ["Layer", "Type", "Output shape", "Parameters"])
    cnn_t += "\n| **Total** | | | **406,369** |"
    toy = list(csv.DictReader(open(ROOT / "reports" / "tuner_toy.csv")))
    toy_t = md_table([(r["trial"], f"{float(r['learning_rate']):.5f}", r["units"], f"{float(r['val_accuracy']):.4f}")
                      for r in toy], ["Trial", "Learning rate", "Units", "Val accuracy"])
    app = list(csv.DictReader(open(ROOT / "reports" / "tuner_visionguard.csv")))
    app_t = md_table([(r["trial"], f"{float(r['fine_tune_lr']):.2e}", r["dropout"], r["unfreeze_layers"],
                       f"{float(r['val_auc']):.4f}") for r in app],
                     ["Trial", "Fine-tune LR", "Dropout", "Unfrozen layers", "Val AUC"])
    return {"CNN_TABLE": cnn_t, "TOY_TABLE": toy_t, "APPLY_TABLE": app_t}


def resolve(text):
    for k, v in tables().items():
        text = text.replace(k, v)
    out = []
    for ln in text.splitlines():
        m = re.fullmatch(r"\[\[OPT (\S+) \| (.+)\]\]", ln.strip())
        if m:
            if (SHOTS / m.group(1)).exists():
                label = Path(m.group(1)).stem
                out.append(f"![label={label}; {m.group(2)}](screenshots/{m.group(1)})")
            continue                                    # missing optional screenshot -> left out
        out.append(ln)
    text = "\n".join(out)
    labels = re.findall(r"!\[label=([\w-]+);", text)
    num = {lab: i + 1 for i, lab in enumerate(labels)}
    missing = set(re.findall(r"\{fig:([\w-]+)\}", text)) - set(num)
    if missing:
        raise SystemExit(f"references to unknown figures: {missing}")
    text = re.sub(r"\{fig:([\w-]+)\}", lambda m: f"Figure {num[m.group(1)]}", text)
    return text, num


def add_figure(doc, n, caption, path):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    with Image.open(ROOT / path) as im:
        w, h = im.size
    width = min(16.0, w * 0.0265 * 1.15)                   # do not blow small screenshots up
    if width * h / w > 19.0:                               # keep tall images on one page
        width = 19.0 * w / h
    p.add_run().add_picture(str(ROOT / path), width=Cm(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    r = cap.add_run(f"Figure {n}: ")
    r.bold = True
    r.font.size = Pt(9)
    add_inline(cap, caption, size=9, color=RGBColor(0x44, 0x4C, 0x5C))


def build():
    text, num = resolve(SRC.read_text(encoding="utf-8"))
    commit = subprocess.check_output(["git", "rev-parse", "--short", "HEAD"], cwd=ROOT, text=True).strip()

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(5)
    for name, size, color in (("Heading 1", 18, "1D2433"), ("Heading 2", 13.5, "2F4A8A")):
        hs = doc.styles[name]
        hs.font.size = Pt(size)
        hs.font.color.rgb = RGBColor.from_string(color)
        hs.font.name = "Calibri"
    footer = sec.footer.paragraphs[0]
    footer.text = "Ali Adnan · 23L0839 · VisionGuard (23L0839_A2_Ops)"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(8)

    title_page(doc, commit)

    lines = text.splitlines()
    i, para, first_h1 = 0, [], True

    def flush():
        if para:
            p = doc.add_paragraph()
            add_inline(p, " ".join(s.strip() for s in para))
            para.clear()

    while i < len(lines):
        s = lines[i].strip()
        if s.startswith("```"):
            flush()
            j, block = i + 1, []
            while j < len(lines) and not lines[j].strip().startswith("```"):
                block.append(lines[j])
                j += 1
            add_code(doc, block)
            i = j + 1
            continue
        if s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue
        m_fig = re.fullmatch(r"!\[label=([\w-]+); (.+)\]\((.+)\)", s)
        if not s:
            flush()
        elif s.startswith("## "):
            flush()
            if not first_h1:
                page_break(doc)
            first_h1 = False
            doc.add_heading(s[3:].replace("`", ""), level=1)
        elif s.startswith("### "):
            flush()
            doc.add_heading(s[4:].replace("`", ""), level=2)
        elif m_fig:
            flush()
            add_figure(doc, num[m_fig.group(1)], m_fig.group(2), m_fig.group(3))
        elif s.startswith("- "):
            flush()
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, s[2:])
        elif re.match(r"^\d+\. ", s):
            flush()
            n, item = s.split(" ", 1)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            add_inline(p, f"**{n}** {item}")
        else:
            if para and re.search(r"[.:!?)`]$", para[-1].strip()):
                flush()
            para.append(s)
        i += 1
    flush()
    doc.save(OUT)
    used = sorted({m for m in re.findall(r"screenshots/(\S+?)\)", text)})
    unused = sorted(p.name for p in SHOTS.glob("*.png") if p.name not in used)
    print(f"saved {OUT}: {len(num)} figures, {len(used)} screenshots used, commit {commit}")
    if unused:
        print("screenshots NOT used:", unused)


if __name__ == "__main__":
    build()
