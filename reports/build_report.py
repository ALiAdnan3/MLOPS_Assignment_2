"""Build the VisionGuard report (Word) from reports/REPORT_NOTES.md.

  python reports/build_report.py   ->  reports/23L0839_VisionGuard.docx

Markers understood in the source text:
  ![caption](path.png)        embedded figure with caption
  [[SHOT id | caption]]       yellow placeholder box where the student pastes a screenshot
Everything else is a small Markdown subset (headings, paragraphs, lists, tables, code blocks, **bold**, `code`).
"""
import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
NOTES = ROOT / "reports" / "REPORT_NOTES.md"
OUT = ROOT / "reports" / "23L0839_VisionGuard.docx"

API_URL = "https://vg-23l0839-a2-ops-api.gentleglacier-7f267617.eastus.azurecontainerapps.io"
LINKS = [
    ("GitHub repository", "https://github.com/ALiAdnan3/MLOPS_Assignment_2"),
    ("DagsHub project (code, DVC data, MLflow, registry)", "https://dagshub.com/ALiAdnan3/23L0839_A2_Ops"),
    ("MLflow tracking server", "https://dagshub.com/ALiAdnan3/23L0839_A2_Ops.mlflow"),
    ("Live model API (Azure Container Apps)", API_URL),
    ("CI toy repository", "https://github.com/ALiAdnan3/23L0839_A2_Ops_ci_toy"),
    ("DVC toy repository", "https://dagshub.com/ALiAdnan3/23L0839_A2_Ops_dvc_toy"),
]

DVC_SECTION = """
### 1.4 DVC – data version control
**Toy part** (`scripts/dvc_toy_lab.ps1`, new empty folder `D:\\dvc_toy_23L0839`):
1. `git init` + `dvc init` – Git versions code and small pointer files; `dvc init` creates `.dvc/` (config + local cache).
2. `students.csv` (2 rows) → `dvc add students.csv` – the file is moved into `.dvc/cache` (named by its MD5), a pointer `students.csv.dvc` is written and `students.csv` is added to `.gitignore`. Commit "v1".
3. Added a row → `dvc status` reports the change → `dvc add` again → new MD5 in the pointer → commit "v2". `git log` shows both versions.
4. **Return to the old version:** `git checkout HEAD~1 -- students.csv.dvc` (old pointer) + `dvc checkout` → the CSV has 2 rows again; then back to v2 the same way.
5. **Remote:** `dvc remote add -d origin s3://dvc` + `endpointurl https://dagshub.com/ALiAdnan3/23L0839_A2_Ops_dvc_toy.s3`; the token is set with `dvc remote modify --local …` so it lands in the git-ignored `.dvc/config.local`. `git push` (both commits to DagsHub) and `dvc push` (data blobs to DagsHub storage).
6. **Delete and pull back:** removed `students.csv` and the whole `.dvc/cache` (a fresh machine) → `dvc pull` → "1 file fetched and 1 file added", content identical.

[[SHOT 2 | Fig 1.4a – Toy: git/dvc init, the pointer file students.csv.dvc and the first commit]]
[[SHOT 3 | Fig 1.4b – Toy: v2 commit, git log with both versions, v1 restored with dvc checkout]]
[[SHOT 4 | Fig 1.4c – Toy: DagsHub remote configured, git push and dvc push]]
[[SHOT 5 | Fig 1.4d – Toy: local data and cache deleted, dvc pull restores the file]]
[[SHOT 6 | Fig 1.4e – Toy repository on DagsHub showing the commits and the DVC-tracked file]]

**Apply:** the real dataset folder `data/raw` (975 files, 480 MB) is tracked the same way: `dvc add data/raw` → `data/raw.dvc` committed to Git, `/raw` added to `data/.gitignore`; remote = DagsHub storage of the main project (`https://dagshub.com/ALiAdnan3/23L0839_A2_Ops.s3`); `dvc push` uploaded all images; `dvc status -c` → "Cache and remote 'origin' are in sync". Anyone can now `git clone` + `dvc pull` the exact dataset. Every training run logs the dataset MD5 as the MLflow tag `dvc_data_version`.

**What the pointer file holds** (`data/raw.dvc`):
```
outs:
- md5: 8f23d6dc9b06d363c3981642c50856c0.dir
  size: 503783378
  nfiles: 975
  hash: md5
  path: raw
```
– the MD5 hash of the content (for a folder: of a manifest listing every file's hash), total size, number of files and the path. It is a tiny text file, so Git can diff and version it; the hash tells DVC exactly which bytes to fetch from cache/remote for any commit.

**Why large data does not belong in Git:** Git keeps every version of every file forever, in every clone – a 480 MB dataset changed a few times makes each clone several GB; Git's delta compression and diff are designed for text, not binary images; GitHub rejects files > 100 MB and warns above 50 MB. DVC keeps the data in cheap object storage, deduplicated by hash, and Git only versions the pointer.

[[SHOT 7 | Fig 1.4f – Apply: data/raw.dvc pointer and `dvc status -c` (cache and DagsHub remote in sync)]]
[[SHOT 14 | Fig 1.4g – Apply: the DVC-tracked data/raw folder in the DagsHub repository]]
"""

CNN_TABLE = """
| Layer | Type | Output shape | Params |
|---|---|---|---|
""" + "\n".join(
    f"| {r['layer']} | {r['type']} | {r['output_shape']} | {int(r['params']):,} |"
    for r in __import__("csv").DictReader(open(ROOT / "reports" / "cnn_layers.csv"))
) + "\n| **Total** | | | **406,369** (405,409 trainable; BN moving statistics are not trainable) |\n"

# (old text, new text) applied in order to the notes
REPLACEMENTS = [
    ("📸 Fig 1.1 – terminal running `python scripts\\check_env.py` (prompt shows user + path).",
     "[[SHOT 1 | Fig 1.1 – Terminal: `python scripts\\check_env.py` – exact versions and the GPU check]]"),
    ("- **Fig 1.3a sample grid** (`reports/figures/sample_grid.png`):", "![Fig 1.3a – Eight random images per class (title = width × height)](reports/figures/sample_grid.png)\n\n- **Fig 1.3a sample grid:**"),
    ("cluster of tiny 10–20 KB thumbnails. File size alone is therefore a weak but non-zero signal – something we do not want the model to rely on.",
     "cluster of tiny 10–20 KB thumbnails. File size alone is therefore a weak but non-zero signal – something we do not want the model to rely on.\n\n"
     "![Fig 1.3b – Width distribution per class](reports/figures/hist_width.png)\n\n"
     "![Fig 1.3c – Height distribution per class](reports/figures/hist_height.png)\n\n"
     "![Fig 1.3d – File-size distribution per class (log scale)](reports/figures/hist_file_size.png)\n" + DVC_SECTION),
    ("📸 Fig 2.1 – `reports/figures/augmented_batch.png`: top row = decoded + resized, bottom = the same images after augmentation.",
     "![Fig 2.1 – One training batch: top = decoded + resized to 224×224, bottom = the same images after random flip / rotation / zoom](reports/figures/augmented_batch.png)"),
    ("📸 Fig 2.2 – `model.summary()` screenshot; table in `reports/cnn_layers.csv`.",
     "Output shape and parameter count per layer:\n" + CNN_TABLE +
     "\n[[SHOT 61 | Fig 2.2a – `model.summary()` of the baseline CNN in the terminal]]"),
    ("ROC-AUC 0.759. 📸 Fig 2.2b `reports/figures/baseline_cnn_training_curves.png`.",
     "ROC-AUC 0.759.\n\n![Fig 2.2b – Baseline CNN training curves (red line = best epoch 16)](reports/figures/baseline_cnn_training_curves.png)"),
    ("📸 Fig 2.3 `reports/figures/mobilenetv2_tl_training_curves.png` – both stages in one figure (dotted line = start of fine-tuning). Trainable-layer lists: `reports/trainable_layers.json`.",
     "![Fig 2.3 – MobileNetV2 training curves: stage 1 (frozen) left of the dotted line, stage 2 (last block fine-tuned) right of it; red line = best epoch 24](reports/figures/mobilenetv2_tl_training_curves.png)\n\n"
     "| Stage | Trainable layers | Trainable params |\n|---|---|---|\n"
     "| 1 – frozen backbone | `ai_probability` (Dense head: kernel + bias) | 1,281 |\n"
     "| 2 – fine-tune last block | `block_16_expand`, `block_16_depthwise`, `block_16_project`, `Conv_1` + `ai_probability` (BatchNorm frozen) | 880,321 |"),
    ("- Fixed seeds do not make CPU/oneDNN kernels",
     "- Fixed seeds do not make CPU/oneDNN kernels"),
    ("Evidence: R1 and R2 have identical settings but val AUC 0.819 vs 0.854.",
     "Evidence: R1 and R2 have identical settings but val AUC 0.819 vs 0.854.\n\n"
     "![Fig 2.4 – Champion (R4) loss and accuracy curves with the best epoch (26) marked](reports/figures/R4_mnv2_no_augment_training_curves.png)"),
    ("📸 Figs: `reports/figures/*_confusion_matrix.png`, `*_roc_curve.png`. ",
     "![Fig 2.5a – Test confusion matrices: baseline CNN, MobileNetV2, champion R4 (each at its validation-chosen threshold)](reports/figures/eval_confusion_matrix_3models.png)\n\n"
     "![Fig 2.5b – Test ROC curves: baseline CNN, MobileNetV2, champion R4](reports/figures/eval_roc_curve_3models.png)\n\n"),
    ("### 2.6 Error analysis (champion, `reports/figures/error_analysis_champion.png`)",
     "### 2.6 Error analysis (champion model R4)\n\n![Fig 2.6 – The 10 most confident test errors of the champion (threshold 0.8)](reports/figures/error_analysis_champion.png)"),
    ("### Bonus – Grad-CAM (`reports/figures/gradcam_champion.png`)",
     "### Bonus – Grad-CAM\n\n![Fig 2.6b – Grad-CAM for 4 correct and 4 wrong test predictions (top = input, bottom = regions driving the predicted class)](reports/figures/gradcam_champion.png)"),
    ("📸 Fig 2.7: terminal output of both tuner searches (tables in `reports/tuner_toy.csv`, `reports/tuner_visionguard.csv`).",
     "(Full tables: `reports/tuner_toy.csv`, `reports/tuner_visionguard.csv`.)"),
    ("📸 Figs 3.1a–d: run list, run page, artifacts, compare view.",
     "[[SHOT 8 | Fig 3.1a – Local MLflow UI: the two toy runs (lr 0.001 and 0.01)]]\n"
     "[[SHOT 9a | Fig 3.1b – Toy run page: parameters, metrics and tags]]\n"
     "[[SHOT 9b | Fig 3.1c – Toy run: per-epoch metric curves]]\n"
     "[[SHOT 9c | Fig 3.1d – Toy run artifacts: loss_curve.png and the logged model]]\n"
     "[[SHOT 10 | Fig 3.1e – MLflow compare view of the two learning rates]]"),
    ("📸 Figs 3.2a–c: DagsHub project page, toy runs on DagsHub, `.env` ignored.",
     "[[SHOT 11 | Fig 3.2a – DagsHub project 23L0839_A2_Ops connected (mirrored) to GitHub]]\n"
     "[[SHOT 12 | Fig 3.2b – The toy runs on the DagsHub MLflow server]]\n"
     "[[SHOT 13 | Fig 3.2c – Proof that .env is git-ignored (`git check-ignore -v .env`, `git status`)]]"),
    ("📸 Figs 3.3a/b: R1 run page (params, metric curves, tags) and R2 autolog run page.",
     "[[SHOT 15 | Fig 3.3a – Manual logging (R1): parameters, metrics and tags git_commit / dvc_data_version / seed]]\n"
     "[[SHOT 16 | Fig 3.3b – Manual logging (R1): per-epoch train/validation curves]]\n"
     "[[SHOT 17 | Fig 3.3c – Autologging (R2): automatically captured parameters (opt_*, callbacks) and tags]]"),
    ("📸 Fig 3.4 artifacts tab of R4.",
     "\n\n[[SHOT 18 | Fig 3.4 – Artifacts of the champion run R4: evaluation plots, split lists, package list, Keras model]]"),
    ("📸 Fig 3.5: DagsHub/MLflow compare view of the runs + our chart.",
     "![Fig 3.5a – Validation metrics used for selection vs test AUC (reported only)](reports/figures/experiment_comparison.png)\n\n"
     "[[SHOT 19 | Fig 3.5b – MLflow compare view of the experiment runs on DagsHub]]"),
    ("📸 Fig 3.6a registry page (versions + aliases), 3.6b terminal output of `predict_champion.py`.",
     "[[SHOT 20 | Fig 3.6a – Model registry: 23L0839_A2_Ops-visionguard versions 1–3 with aliases @champion and @challenger]]\n"
     "[[SHOT 21 | Fig 3.6b – `scripts/predict_champion.py`: model loaded by alias and one image classified]]"),
    ("📸 Fig 4.1 terminal output.", "[[SHOT 22 | Fig 4.1 – Inference module console output for two known images]]"),
    ("📸 Figs 4.2a–c: the three failure pages (wrong type, too large, undecodable).",
     "[[SHOT 30 | Fig 4.2a – Wrong file type rejected (415)]]\n"
     "[[SHOT 31 | Fig 4.2b – File over the size limit rejected (413)]]\n"
     "[[SHOT 32 | Fig 4.2c – File with an image extension that cannot be decoded (422)]]"),
    ("📸 Figs 4.3a–c: upload page, a result page, history page.",
     "[[SHOT 27 | Fig 4.3a – Upload page]]\n"
     "[[SHOT 28 | Fig 4.3b – Result page: AI-generated image with confidence and caveat]]\n"
     "[[SHOT 29 | Fig 4.3c – Result page: real image]]"),
    ("📸 Fig 4.4 `/health` in the browser.", "[[SHOT 34 | Fig 4.4 – /health response (debug false, database ok)]]"),
    ("📸 Fig 5.1 History page; Fig 5.2 psql query.",
     "[[SHOT 33 | Fig 5.1 – History menu: every prediction from the database]]\n"
     "(The same rows queried with psql inside the db container are shown in Fig 6.5b.)"),
    ("📸 Figs 6.1a–d (one per pause of the script).",
     "[[SHOT 35 | Fig 6.1a – docker run with port mapping and volume mount, docker ps, page served from the mounted folder]]\n"
     "[[SHOT 36 | Fig 6.1b – docker logs and a shell inside the container (docker exec)]]\n"
     "[[SHOT 37 | Fig 6.1c – Listing and removing the container and the image]]"),
    ("📸 Fig 6.2 `Dockerfile` on GitHub + the `id` output.",
     "Dockerfile: https://github.com/ALiAdnan3/MLOPS_Assignment_2/blob/main/Dockerfile\n\n"
     "[[SHOT 44a | Fig 6.2 – `docker run --rm --entrypoint id 23l0839_a2_ops-web` → uid=10001(appuser)]]"),
    ("📸 Figs 6.3a compose file on GitHub, 6.3b `docker compose ps` (both healthy), 6.3c port 5432 closed.",
     "Compose file: https://github.com/ALiAdnan3/MLOPS_Assignment_2/blob/main/docker-compose.yml\n\n"
     "[[SHOT 44b | Fig 6.3 – `docker compose ps` (both services healthy, db shows 5432/tcp with no host mapping) and `Test-NetConnection 127.0.0.1 -Port 5432` → False]]"),
    ("📸 Fig 6.4 `docker images` with both tags + `docker history`.",
     "[[SHOT 62 | Fig 6.4 – `docker images 23l0839_a2_ops-web` (before/after) and `docker history`]]"),
    ("📸 Figs 6.5a–e (one per pause of the script).",
     "[[SHOT 39 | Fig 6.5a – start, stop, start, restart and status of both services]]\n"
     "[[SHOT 40 | Fig 6.5b – logs of both services, shell in web (appuser) and psql in db]]\n"
     "[[SHOT 41 | Fig 6.5c – Failure: web restarting, log ends with 'password authentication failed']]\n"
     "[[SHOT 42 | Fig 6.5d – Diagnosis from the web and db logs]]\n"
     "[[SHOT 43 | Fig 6.5e – Fixed: both services healthy, /health database ok]]"),
    ("📸 Figs 6.6a toy runs list (✅❌✅), 6.6b red job log, 6.6c green matrix, 6.6d project CI run, 6.6e badge in README.",
     "[[SHOT 24 | Fig 6.6a – Toy repository Actions: green → red → green]]\n"
     "[[SHOT 25 | Fig 6.6b – Red run: failing pytest step]]\n"
     "[[SHOT 26 | Fig 6.6c – Green run: both matrix jobs (Python 3.11 and 3.12) pass]]\n"
     "[[SHOT 23 | Fig 6.6d – VisionGuard CI: lint-and-test and docker-build jobs green]]\n"
     "[[SHOT 63 | Fig 6.6e – CI status badge in the README on GitHub]]"),
    ("📸 Figs 7.1 resource group in Azure portal, 7.2 ACR repository with tag v2, 7.3 Container App overview (URL, running), 7.4 env vars/secrets page (values hidden), 7.5 log stream.",
     "[[SHOT 45 | Fig 7.1 – Azure resource group 23L0839_A2_Ops with all resources]]\n"
     "[[SHOT 46 | Fig 7.2 – Azure Container Registry: repository 23l0839-a2-ops-api, tag v2]]\n"
     "[[SHOT 47 | Fig 7.3 – Container App overview: application URL, status Running]]\n"
     "[[SHOT 48 | Fig 7.4 – Container App environment variables: API_KEY as secret reference, THRESHOLD, MODEL_VERSION]]\n"
     "[[SHOT 49 | Fig 7.5 – Container App log stream showing predict requests]]"),
    ("📸 Figs 7.6a–d Postman (each request + test results). **Tested from another network:** open the `/health` URL on a phone using mobile data (Wi-Fi off) – 📸 Fig 7.7.",
     "[[SHOT 50 | Fig 7.6a – Postman: GET /health]]\n"
     "[[SHOT 51 | Fig 7.6b – Postman: POST /predict with an AI image (tests passed)]]\n"
     "[[SHOT 52 | Fig 7.6c – Postman: POST /predict with a real image]]\n"
     "[[SHOT 53 | Fig 7.6d – Postman: POST /predict without API key → 401]]\n\n"
     "**Tested from another network:** the `/health` URL opened on a phone over mobile data (Wi-Fi off):\n\n"
     "[[SHOT 54 | Fig 7.7 – API reached from a mobile network]]"),
    ("📸 Figs 8.1–8.4 several result pages in remote mode (model shown as `remote`), 8.5 History page, 8.6 `/health` showing `inference_mode: remote` and the `model_api` block, 8.7 Azure log stream showing the incoming `predict` requests.",
     "[[SHOT 55 | Fig 8.1 – Result page in remote mode (AI image #1)]]\n"
     "[[SHOT 56 | Fig 8.2 – Result page in remote mode (AI image #2)]]\n"
     "[[SHOT 57 | Fig 8.3 – Result page in remote mode (real image #1)]]\n"
     "[[SHOT 58 | Fig 8.4 – Result page in remote mode (real image #2)]]\n"
     "[[SHOT 59 | Fig 8.5 – History menu with local and remote rows from the database]]\n"
     "[[SHOT 60 | Fig 8.6 – /health in remote mode including the cloud model_api block]]\n"
     "(The matching requests arriving in Azure are visible in the log stream, Fig 7.5.)"),
]


# ----------------------------------------------------------------------------- docx helpers
def shade(cell_or_par, hex_color):
    el = cell_or_par._tc if hasattr(cell_or_par, "_tc") else cell_or_par._p
    pr = el.get_or_add_tcPr() if hasattr(cell_or_par, "_tc") else el.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    pr.append(shd)


def border(par, color="C9A227", size=12):
    pPr = par._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "dashed")
        e.set(qn("w:sz"), str(size))
        e.set(qn("w:space"), "6")
        e.set(qn("w:color"), color)
        b.append(e)
    pPr.append(b)


INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)")


def add_inline(par, text, size=None, color=None):
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            run = par.add_run(tok[2:-2])
            run.bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            run = par.add_run(tok[1:-1])
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.color.rgb = RGBColor(0x9C, 0x1F, 0x4F)
            run.font.size = Pt((size or 10.5) - 1)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            run = par.add_run(tok[1:-1])
            run.italic = True
        else:
            run = par.add_run(tok)
        if size and not (tok.startswith("`")):
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.paragraphs[0].text = ""
        add_inline(c.paragraphs[0], h, size=9)
        for r in c.paragraphs[0].runs:
            r.bold = True
        shade(c, "DCE6F2")
    for row in body:
        cells = t.add_row().cells
        for i in range(len(header)):
            val = row[i] if i < len(row) else ""
            cells[i].paragraphs[0].text = ""
            add_inline(cells[i].paragraphs[0], val, size=9)
    doc.add_paragraph()


def add_figure(doc, caption, path):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ROOT / path), width=Cm(16))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(cap, caption, size=9, color=RGBColor(0x44, 0x4C, 0x5C))
    for r in cap.runs:
        r.italic = True


def add_shot(doc, sid, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    shade(p, "FFF6D5")
    border(p)
    r = p.add_run(f"📸  SCREENSHOT {sid} – paste here")
    r.bold = True
    r.font.color.rgb = RGBColor(0x8A, 0x61, 0x00)
    p.add_run("\n")
    add_inline(p, "(Click this box, paste the screenshot with Ctrl+V, then delete this text. "
                  "Show your username / address bar and the clock.)", size=8.5, color=RGBColor(0x8A, 0x61, 0x00))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(cap, caption, size=9, color=RGBColor(0x44, 0x4C, 0x5C))
    for run in cap.runs:
        run.italic = True


def add_code(doc, lines):
    for i, ln in enumerate(lines):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        shade(p, "F2F4F7")
        r = p.add_run(ln if ln else " ")
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        r.font.size = Pt(8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def toc(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    for tag, text in (("begin", None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None), (None, None),
                      ("end", None)):
        if tag:
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), tag)
            run._r.append(f)
        elif text:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            run._r.append(it)
        else:
            t = OxmlElement("w:t")
            t.text = "Right-click → Update Field to build the table of contents."
            run._r.append(t)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ----------------------------------------------------------------------------- main
def title_page(doc, commit):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("VisionGuard")
    r.bold = True
    r.font.size = Pt(34)
    r.font.color.rgb = RGBColor(0x1D, 0x24, 0x33)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("AI vs Real Image Classification and Deployment")
    r.font.size = Pt(16)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("MLOps – Assignment 2 · Project 23L0839_A2_Ops")
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x6B, 0x73, 0x85)
    doc.add_paragraph()
    rows = [("Name", "Ali Adnan"), ("Roll number", "23L0839"), ("Project name", "23L0839_A2_Ops")]
    rows += LINKS
    rows += [("Graded commit", commit), ("Champion model", "23L0839_A2_Ops-visionguard v2 (@champion) – MobileNetV2, threshold 0.80")]
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, v in rows:
        c = t.add_row().cells
        c[0].text = k
        c[0].paragraphs[0].runs[0].bold = True
        c[0].paragraphs[0].runs[0].font.size = Pt(9.5)
        shade(c[0], "EEF2F8")
        c[1].text = v
        c[1].paragraphs[0].runs[0].font.size = Pt(9.5)
        c[0].width, c[1].width = Cm(5.2), Cm(11.3)
    doc.add_paragraph()
    p = doc.add_paragraph()
    add_inline(p, "**Disclosure of AI assistance:** Claude Code (Anthropic) was used as a tutor and pair-programmer: to explain "
                  "the concepts of each task, to scaffold and debug the code (training pipeline, Flask app, Dockerfiles, "
                  "workflows, deployment commands) and to draft the explanatory text of this report. I ran every command, "
                  "checked every result and took all screenshots myself. Official documentation of TensorFlow/Keras, "
                  "MLflow, DagsHub, DVC, Docker and Azure was consulted for the tool labs.", size=9.5)
    page_break(doc)
    h = doc.add_paragraph()
    r = h.add_run("Contents")
    r.bold = True
    r.font.size = Pt(16)
    toc(doc)
    page_break(doc)


def build():
    text = NOTES.read_text(encoding="utf-8")
    # drop the draft header (replaced by the title page)
    text = text.split("\n## Task 1", 1)[1]
    text = "## Task 1" + text
    for old, new in REPLACEMENTS:
        if old not in text:
            raise SystemExit(f"replacement anchor not found: {old[:70]}")
        text = text.replace(old, new, 1)
    before = (ROOT / "outputs" / "before_size.txt")
    text = text.replace("BEFORE_SIZE", before.read_text().strip() if before.exists() else "see Fig 6.4")
    left = [ln for ln in text.splitlines() if "📸" in ln and "[[SHOT" not in ln]
    if left:
        raise SystemExit("unconverted screenshot references:\n" + "\n".join(left))

    commit = subprocess.check_output(["git", "rev-parse", "--short", "HEAD"], cwd=ROOT, text=True).strip()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(4)
    for name, size, color in (("Heading 1", 18, "1D2433"), ("Heading 2", 13, "2F4A8A")):
        hs = doc.styles[name]
        hs.font.size = Pt(size)
        hs.font.color.rgb = RGBColor.from_string(color)
        hs.font.name = "Calibri"
    footer = sec.footer.paragraphs[0]
    footer.text = "23L0839 · Ali Adnan · VisionGuard (23L0839_A2_Ops)"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(8)

    title_page(doc, commit)

    lines = text.splitlines()
    i, para, first_h1 = 0, [], True

    def flush():
        if para:
            p = doc.add_paragraph()
            add_inline(p, " ".join(s.strip() for s in para))
            para.clear()

    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if s.startswith("```"):
            flush()
            j = i + 1
            block = []
            while j < len(lines) and not lines[j].strip().startswith("```"):
                block.append(lines[j])
                j += 1
            add_code(doc, block)
            i = j + 1
            continue
        if s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue
        m_fig = re.fullmatch(r"!\[(.+)\]\((.+)\)", s)
        m_shot = re.fullmatch(r"\[\[SHOT ([\w]+) \| (.+)\]\]", s)
        if not s or s == "---":
            flush()
        elif s.startswith("## "):
            flush()
            if not first_h1:
                page_break(doc)
            first_h1 = False
            doc.add_heading(s[3:].replace("`", ""), level=1)
        elif s.startswith("### "):
            flush()
            doc.add_heading(s[4:].replace("`", ""), level=2)
        elif m_fig:
            flush()
            add_figure(doc, m_fig.group(1), m_fig.group(2))
        elif m_shot:
            flush()
            add_shot(doc, m_shot.group(1), m_shot.group(2))
        elif re.match(r"^(\s*)- ", ln):
            flush()
            indent = len(re.match(r"^(\s*)", ln).group(1))
            item = s[2:]
            while i + 1 < len(lines) and lines[i + 1].startswith("  ") and not lines[i + 1].strip().startswith(("-", "|")) \
                    and lines[i + 1].strip():
                i += 1
                item += " " + lines[i].strip()
            p = doc.add_paragraph(style="List Bullet" if indent == 0 else "List Bullet 2")
            add_inline(p, item)
        elif re.match(r"^\d+\. ", s):
            flush()
            num, item = s.split(" ", 1)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            add_inline(p, f"**{num}** {item}")
        else:
            # join wrapped lines, but start a new paragraph after a finished sentence
            if para and re.search(r"[.:!?)`]$", para[-1].strip()):
                flush()
            para.append(s)
        i += 1
    flush()
    doc.save(OUT)
    shots = re.findall(r"\[\[SHOT ([\w]+) \|", text)
    print(f"saved {OUT} – {len(shots)} screenshot boxes, commit {commit}")


if __name__ == "__main__":
    build()
